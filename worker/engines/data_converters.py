"""Hardened data and document format converters for Phase 8.

All processing occurs within disposable sandboxed containers under non-root execution.
"""

import csv
import io
import json
import os
import re
import subprocess
from pathlib import Path

try:
    from validation import MAX_INPUT, MAX_OUTPUT, validate_office, validate_pdf
except ImportError:
    from engines.validation import MAX_INPUT, MAX_OUTPUT, validate_office, validate_pdf


FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")


def get_work_dir() -> Path:
    target = Path(os.environ.get("CONVERTER_WORK_DIR", "/work"))
    try:
        if target.exists() and os.access(target, os.W_OK):
            return target
    except Exception:  # noqa: S110
        pass
    return Path("/tmp")  # noqa: S108 -- fallback to private container tmpfs


def escape_formula(value: str) -> str:
    """Neutralize spreadsheet formula injection."""
    if value and value[0] in FORMULA_PREFIXES:
        return "'" + value
    return value


def csv_to_xlsx(data: bytes):
    from openpyxl import Workbook, load_workbook

    if len(data) > MAX_INPUT:
        raise ValueError("RESOURCE_LIMIT")
    text = data.decode("utf-8-sig")
    if "\x00" in text:
        raise ValueError("CORRUPT_FILE")

    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except Exception:
        dialect = csv.excel

    reader = csv.reader(io.StringIO(text), dialect=dialect)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"

    row_count = col_count = 0
    for row_idx, row in enumerate(reader, 1):
        row_count += 1
        if row_count > 50000:
            raise ValueError("RESOURCE_LIMIT")
        if len(row) > 1000:
            raise ValueError("RESOURCE_LIMIT")
        col_count = max(col_count, len(row))
        for col_idx, cell_value in enumerate(row, 1):
            cell = sheet.cell(row_idx, col_idx)
            val = escape_formula(cell_value)
            cell.value = val
            cell.data_type = "s"

    output = get_work_dir() / "output.xlsx"
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)
    workbook.close()

    if output.stat().st_size > MAX_OUTPUT:
        raise ValueError("RESOURCE_LIMIT")

    verified = load_workbook(output, read_only=True, data_only=False)
    try:
        if any(cell.data_type == "f" for s in verified for r in s for cell in r):
            raise ValueError("INVALID_OUTPUT")
    finally:
        verified.close()

    return output, {
        "pages": 1,
        "rows": row_count,
        "columns": col_count,
        "mediaType": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    }


def xlsx_to_csv(data: bytes):
    from openpyxl import load_workbook

    validate_office(data, "xlsx")
    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        sheet = workbook.active
        output = get_work_dir() / "output.csv"
        output.parent.mkdir(parents=True, exist_ok=True)
        with output.open("w", newline="", encoding="utf-8") as stream:
            writer = csv.writer(stream)
            row_count = 0
            for row in sheet.iter_rows(values_only=True):
                row_count += 1
                if row_count > 50000:
                    raise ValueError("RESOURCE_LIMIT")
                cleaned_row = [escape_formula(str(val)) if val is not None else "" for val in row]
                writer.writerow(cleaned_row)
    finally:
        workbook.close()

    if output.stat().st_size > MAX_OUTPUT:
        raise ValueError("RESOURCE_LIMIT")

    return output, {
        "pages": 1,
        "rows": row_count,
        "mediaType": "text/csv; charset=utf-8",
    }


def json_to_xlsx(data: bytes):
    from openpyxl import Workbook, load_workbook

    if len(data) > MAX_INPUT:
        raise ValueError("RESOURCE_LIMIT")
    try:
        parsed = json.loads(data.decode("utf-8-sig"))
    except Exception:
        raise ValueError("CORRUPT_FILE") from None

    if not isinstance(parsed, (list, dict)):
        raise ValueError("UNSUPPORTED_FORMAT")

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "JSON_Data"

    if isinstance(parsed, dict):
        parsed = [parsed]

    if not parsed:
        sheet.cell(1, 1).value = ""
    elif isinstance(parsed[0], list):
        for row_idx, row in enumerate(parsed, 1):
            if row_idx > 50000 or len(row) > 1000:
                raise ValueError("RESOURCE_LIMIT")
            for col_idx, val in enumerate(row, 1):
                cell = sheet.cell(row_idx, col_idx)
                cell.value = escape_formula(str(val) if val is not None else "")
                cell.data_type = "s"
    elif isinstance(parsed[0], dict):
        headers = []
        for item in parsed:
            if isinstance(item, dict):
                for k in item.keys():
                    if k not in headers:
                        headers.append(k)
        if len(headers) > 1000:
            raise ValueError("RESOURCE_LIMIT")

        for col_idx, h in enumerate(headers, 1):
            cell = sheet.cell(1, col_idx)
            cell.value = str(h)
            cell.data_type = "s"

        for row_idx, item in enumerate(parsed, 2):
            if row_idx > 50000:
                raise ValueError("RESOURCE_LIMIT")
            if isinstance(item, dict):
                for col_idx, h in enumerate(headers, 1):
                    val = item.get(h)
                    cell = sheet.cell(row_idx, col_idx)
                    cell.value = escape_formula(
                        json.dumps(val)
                        if isinstance(val, (dict, list))
                        else str(val)
                        if val is not None
                        else ""
                    )
                    cell.data_type = "s"

    output = get_work_dir() / "output.xlsx"
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)
    workbook.close()

    if output.stat().st_size > MAX_OUTPUT:
        raise ValueError("RESOURCE_LIMIT")

    verified = load_workbook(output, read_only=True)
    verified.close()

    return output, {
        "pages": 1,
        "mediaType": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    }


def xlsx_to_json(data: bytes):
    from openpyxl import load_workbook

    validate_office(data, "xlsx")
    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        sheet = workbook.active
        rows = list(sheet.iter_rows(values_only=True))
    finally:
        workbook.close()

    if not rows:
        result = []
    else:
        headers = [
            str(cell) if cell is not None else f"col_{idx}" for idx, cell in enumerate(rows[0], 1)
        ]
        result = []
        for row in rows[1:]:
            if len(result) > 50000:
                raise ValueError("RESOURCE_LIMIT")
            record = {}
            for col_idx, header in enumerate(headers):
                val = row[col_idx] if col_idx < len(row) else None
                record[header] = val
            result.append(record)

    output = get_work_dir() / "output.json"
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(result, indent=2, ensure_ascii=False), encoding="utf-8")
    return output, {
        "pages": 1,
        "count": len(result),
        "mediaType": "application/json; charset=utf-8",
    }


def markdown_to_docx(data: bytes):
    from docx import Document

    if len(data) > MAX_INPUT:
        raise ValueError("RESOURCE_LIMIT")
    text = data.decode("utf-8-sig")
    if "\x00" in text:
        raise ValueError("CORRUPT_FILE")

    doc = Document()
    lines = text.splitlines()
    in_code_block = False
    code_lines = []

    for line in lines:
        if line.startswith("```"):
            if in_code_block:
                p = doc.add_paragraph("\n".join(code_lines))
                p.style = "Normal"
                code_lines = []
                in_code_block = False
            else:
                in_code_block = True
            continue

        if in_code_block:
            code_lines.append(line)
            continue

        stripped = line.strip()
        if not stripped:
            continue

        if stripped.startswith("# "):
            doc.add_heading(stripped[2:], level=1)
        elif stripped.startswith("## "):
            doc.add_heading(stripped[3:], level=2)
        elif stripped.startswith("### "):
            doc.add_heading(stripped[4:], level=3)
        elif stripped.startswith("#### "):
            doc.add_heading(stripped[5:], level=4)
        elif stripped.startswith("- ") or stripped.startswith("* "):
            doc.add_paragraph(stripped[2:], style="List Bullet")
        elif re.match(r"^\d+\.\s+", stripped):
            content = re.sub(r"^\d+\.\s+", "", stripped)
            doc.add_paragraph(content, style="List Number")
        else:
            p = doc.add_paragraph()
            # Simple markdown bold parsing
            parts = re.split(r"(\*\*.*?\*\*)", line)
            for part in parts:
                if part.startswith("**") and part.endswith("**") and len(part) >= 4:
                    p.add_run(part[2:-2]).bold = True
                else:
                    p.add_run(part)

    output = get_work_dir() / "output.docx"
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)

    if output.stat().st_size > MAX_OUTPUT:
        raise ValueError("RESOURCE_LIMIT")

    return output, {
        "pages": 1,
        "mediaType": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    }


def docx_to_markdown(data: bytes):
    from docx import Document

    validate_office(data, "docx")
    doc = Document(io.BytesIO(data))
    md_lines = []

    for p in doc.paragraphs:
        text = ""
        for run in p.runs:
            run_text = run.text
            if run.bold and run.italic:
                run_text = f"***{run_text}***"
            elif run.bold:
                run_text = f"**{run_text}**"
            elif run.italic:
                run_text = f"*{run_text}*"
            text += run_text

        text = text.strip()
        if not text:
            continue

        style_name = p.style.name.lower()
        if "heading 1" in style_name:
            md_lines.append(f"# {text}")
        elif "heading 2" in style_name:
            md_lines.append(f"## {text}")
        elif "heading 3" in style_name:
            md_lines.append(f"### {text}")
        elif "bullet" in style_name:
            md_lines.append(f"- {text}")
        elif "number" in style_name:
            md_lines.append(f"1. {text}")
        else:
            md_lines.append(text)
        md_lines.append("")

    for table in doc.tables:
        t_rows = []
        for row in table.rows:
            t_rows.append([cell.text.strip().replace("\n", " ") for cell in row.cells])
        if t_rows:
            header = t_rows[0]
            md_lines.append("| " + " | ".join(header) + " |")
            md_lines.append("| " + " | ".join(["---"] * len(header)) + " |")
            for r in t_rows[1:]:
                md_lines.append("| " + " | ".join(r) + " |")
            md_lines.append("")

    output = get_work_dir() / "output.md"
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text("\n".join(md_lines).strip() + "\n", encoding="utf-8")
    return output, {
        "pages": 1,
        "mediaType": "text/markdown; charset=utf-8",
    }


def markdown_to_pdf(data: bytes):
    from weasyprint import HTML

    if len(data) > 2 * 1024 * 1024:
        raise ValueError("RESOURCE_LIMIT")
    text = data.decode("utf-8-sig")
    if "\x00" in text:
        raise ValueError("CORRUPT_FILE")

    # Format text paragraphs to HTML
    lines = text.splitlines()
    body_parts = []
    for line in lines:
        s = line.strip()
        if s.startswith("# "):
            body_parts.append(f"<h1>{s[2:]}</h1>")
        elif s.startswith("## "):
            body_parts.append(f"<h2>{s[3:]}</h2>")
        elif s.startswith("### "):
            body_parts.append(f"<h3>{s[4:]}</h3>")
        elif s.startswith("- "):
            body_parts.append(f"<li>{s[2:]}</li>")
        elif s:
            body_parts.append(f"<p>{s}</p>")

    html_content = f"""<!DOCTYPE html>
<html>
<head>
<meta charset="utf-8">
<style>
  body {{ font-family: Liberation Sans, sans-serif; margin: 20mm; font-size: 11pt; line-height: 1.5; }}
  h1 {{ font-size: 20pt; margin-bottom: 8pt; }}
  h2 {{ font-size: 15pt; margin-top: 12pt; margin-bottom: 6pt; }}
  h3 {{ font-size: 12pt; margin-top: 10pt; margin-bottom: 4pt; }}
  p {{ margin-bottom: 8pt; }}
  li {{ margin-left: 15pt; }}
</style>
</head>
<body>
{"".join(body_parts)}
</body>
</html>"""

    blocked = []

    def reject_resource(url, *args, **kwargs):
        blocked.append(True)
        raise ValueError("External resources are disabled")

    output = get_work_dir() / "output.pdf"
    output.parent.mkdir(parents=True, exist_ok=True)
    doc = HTML(string=html_content, url_fetcher=reject_resource).render()
    if blocked or not 1 <= len(doc.pages) <= 100:
        raise ValueError("UNSUPPORTED_FORMAT")
    doc.write_pdf(output)
    pages = validate_pdf(output)
    return output, {"pages": pages, "mediaType": "application/pdf"}


def txt_to_pdf(data: bytes):
    from weasyprint import HTML

    if len(data) > 2 * 1024 * 1024:
        raise ValueError("RESOURCE_LIMIT")
    text = data.decode("utf-8-sig")
    if "\x00" in text:
        raise ValueError("CORRUPT_FILE")

    # Escape HTML special chars
    escaped = text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")

    html_content = f"""<!DOCTYPE html>
<html>
<head>
<meta charset="utf-8">
<style>
  body {{
    font-family: monospace;
    font-size: 10pt;
    line-height: 1.4;
    white-space: pre-wrap;
    word-break: break-all;
    margin: 15mm;
  }}
</style>
</head>
<body>{escaped}</body>
</html>"""

    blocked = []

    def reject_resource(url, *args, **kwargs):
        blocked.append(True)
        raise ValueError("External resources are disabled")

    output = get_work_dir() / "output.pdf"
    output.parent.mkdir(parents=True, exist_ok=True)
    doc = HTML(string=html_content, url_fetcher=reject_resource).render()
    if blocked or not 1 <= len(doc.pages) <= 100:
        raise ValueError("UNSUPPORTED_FORMAT")
    doc.write_pdf(output)
    pages = validate_pdf(output)
    return output, {"pages": pages, "mediaType": "application/pdf"}


def rtf_to_pdf(data: bytes):
    if len(data) > MAX_INPUT:
        raise ValueError("RESOURCE_LIMIT")
    if not data.startswith(b"{\\rtf"):
        raise ValueError("CORRUPT_FILE")

    source = Path("/work/input.rtf")
    source.write_bytes(data)
    profile = Path("/work/profile/user")
    profile.mkdir(parents=True, exist_ok=True)
    (profile / "registrymodifications.xcu").write_text("""<?xml version="1.0" encoding="UTF-8"?>
<oor:items xmlns:oor="http://openoffice.org/2001/registry"><item oor:path="/org.openoffice.Office.Common/Security/Scripting"><prop oor:name="MacroSecurityLevel" oor:op="fuse"><value>3</value></prop></item><item oor:path="/org.openoffice.Office.Common/Security/Scripting"><prop oor:name="DisableMacrosExecution" oor:op="fuse"><value>true</value></prop></item></oor:items>""")
    subprocess.run(  # noqa: S603 -- fixed executable and generated sandbox paths; no shell.
        [
            "/usr/bin/libreoffice",
            "-env:UserInstallation=file:///work/profile",
            "--headless",
            "--nologo",
            "--nodefault",
            "--norestore",
            "--convert-to",
            "pdf",
            "--outdir",
            "/work",
            str(source),
        ],
        stdout=subprocess.DEVNULL,
        stderr=subprocess.DEVNULL,
        timeout=90,
        check=True,
    )
    output = Path("/work/input.pdf")
    pages = validate_pdf(output)
    return output, {"pages": pages, "mediaType": "application/pdf"}
