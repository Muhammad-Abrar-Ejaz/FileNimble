"""Explicitly limited reverse conversions: slide images and machine-generated tables."""

import io
from pathlib import Path

from validation import MAX_OUTPUT, validate_pdf_input


def convert(data, tool):
    source = Path("/work/input.pdf")
    source.write_bytes(data)
    pages = validate_pdf_input(source)
    if pages > 40:
        raise ValueError("RESOURCE_LIMIT")
    if tool == "pdf-to-powerpoint":
        return slides(source, pages)
    return tables(source, pages)


def slides(source, pages):
    import pypdfium2 as pdfium
    from pptx import Presentation
    from pptx.util import Inches

    output = Path("/work/output.pptx")
    presentation = Presentation()
    presentation.slide_width, presentation.slide_height = Inches(10), Inches(7.5)
    total = 0
    with pdfium.PdfDocument(source) as document:
        for index in range(pages):
            page = document[index]
            try:
                width, height = page.get_size()
                scale = 96 / 72
                pixels = int(width * scale + 1) * int(height * scale + 1)
                total += pixels
                if pixels > 4_000_000 or total > 64_000_000:
                    raise ValueError("RESOURCE_LIMIT")
                bitmap = page.render(scale=scale, draw_annots=False)
                try:
                    image = bitmap.to_pil()
                    buffer = io.BytesIO()
                    image.save(buffer, format="PNG")
                    image.close()
                    buffer.seek(0)
                    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
                    ratio = min(
                        presentation.slide_width / width, presentation.slide_height / height
                    )
                    draw_width, draw_height = int(width * ratio), int(height * ratio)
                    slide.shapes.add_picture(
                        buffer,
                        (presentation.slide_width - draw_width) // 2,
                        (presentation.slide_height - draw_height) // 2,
                        width=draw_width,
                        height=draw_height,
                    )
                    buffer.close()
                finally:
                    bitmap.close()
            finally:
                page.close()
    presentation.save(output)
    if output.stat().st_size > MAX_OUTPUT:
        raise ValueError("RESOURCE_LIMIT")
    verified = Presentation(output)
    if len(verified.slides) != pages or any(len(slide.shapes) != 1 for slide in verified.slides):
        raise ValueError("INVALID_OUTPUT")
    return output, {
        "pages": pages,
        "mediaType": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    }


def tables(source, pages):
    import pdfplumber
    from openpyxl import Workbook, load_workbook

    output = Path("/work/output.xlsx")
    workbook = Workbook()
    workbook.remove(workbook.active)
    count = cells = rows = 0
    with pdfplumber.open(source) as document:
        for page_number, page in enumerate(document.pages, 1):
            for table in page.extract_tables():
                if len(table) < 2 or max((len(row) for row in table), default=0) < 2:
                    continue
                count += 1
                if count > 100:
                    raise ValueError("RESOURCE_LIMIT")
                sheet = workbook.create_sheet(f"Page {page_number} Table {count}")
                for row_number, row in enumerate(table, 1):
                    rows += 1
                    cells += len(row)
                    if rows > 10000 or cells > 100000 or len(row) > 100:
                        raise ValueError("RESOURCE_LIMIT")
                    for column, value in enumerate(row, 1):
                        text = value or ""
                        if len(text) > 4096:
                            raise ValueError("RESOURCE_LIMIT")
                        cell = sheet.cell(row_number, column)
                        cell.value = text
                        cell.data_type = (
                            "s"  # Extracted content is text, never executable formulas.
                        )
            page.close()
    if not count:
        raise ValueError("UNSUPPORTED_FORMAT")
    workbook.save(output)
    workbook.close()
    if output.stat().st_size > MAX_OUTPUT:
        raise ValueError("RESOURCE_LIMIT")
    verified = load_workbook(output, read_only=True, data_only=False)
    try:
        if len(verified.sheetnames) != count or any(
            cell.data_type == "f" for sheet in verified for row in sheet for cell in row
        ):
            raise ValueError("INVALID_OUTPUT")
    finally:
        verified.close()
    return output, {
        "pages": pages,
        "mediaType": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    }
