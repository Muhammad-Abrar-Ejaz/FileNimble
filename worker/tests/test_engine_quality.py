import io

import pytest
from pypdf import PdfReader

from native_fixtures import HTML, office_fixture
from qualify_office import qualify
from app.config import settings

pytestmark = pytest.mark.skipif(
    not settings.native_test, reason="Native qualification requires compose.native-test.yaml"
)


@pytest.mark.parametrize(
    "tool,extension,words",
    [
        ("word-to-pdf", "docx", ("First page", "Second page")),
        ("powerpoint-to-pdf", "pptx", ("First slide", "Second slide")),
        ("excel-to-pdf", "xlsx", ("First sheet", "Second sheet")),
    ],
)
def test_office_page_content(tool, extension, words):
    result = qualify("512m", tool, office_fixture(extension), expected_pages=2)
    reader = PdfReader(io.BytesIO(result), strict=True)
    assert len(reader.pages) == 2
    for page, word in zip(reader.pages, words, strict=True):
        assert "".join(word.split()) in "".join(page.extract_text().split())


def test_html_and_table_text():
    result = qualify("512m", "html-to-pdf", HTML, "html")
    text = PdfReader(io.BytesIO(result)).pages[0].extract_text()
    assert "Owned HTML fixture" in text and "Alpha" in text and "34" in text


@pytest.mark.parametrize(
    "url",
    [
        "http://127.0.0.1/private",
        "http://169.254.169.254/",
        "http://10.0.0.1/",
        "https://example.com/a",
        "file:///etc/passwd",
        "data:image/png;base64,eA==",
    ],
)
def test_html_resource_requests_rejected(url):
    qualify(
        "512m",
        "html-to-pdf",
        f'<html><body><img src="{url}"></body></html>'.encode(),
        "html",
        expect_error=True,
    )


def test_pdf_utilities_preserve_content():
    import re

    source = qualify("512m", "html-to-pdf", HTML, "html")
    compressed = qualify("256m", "compress-pdf", source, "utility")
    assert "Alpha" in PdfReader(io.BytesIO(compressed)).pages[0].extract_text()
    damaged = re.sub(rb"startxref\s+\d+", b"startxref\n0", source)
    repaired = qualify("256m", "repair-pdf", damaged, "utility")
    assert "Alpha" in PdfReader(io.BytesIO(repaired), strict=True).pages[0].extract_text()


def test_image_based_slides_and_table_cells():
    from pptx import Presentation
    from openpyxl import load_workbook

    source = qualify("512m", "html-to-pdf", HTML, "html")
    slides = qualify("512m", "pdf-to-powerpoint", source, "extraction")
    presentation = Presentation(io.BytesIO(slides))
    assert len(presentation.slides) == 1
    shape = presentation.slides[0].shapes[0]
    assert shape.shape_type == 13  # Picture, not editable reconstructed content.
    assert shape.image.blob.startswith(b"\x89PNG")
    tables = qualify("512m", "pdf-to-excel", source, "extraction")
    workbook = load_workbook(io.BytesIO(tables), data_only=False)
    try:
        values = [cell.value for sheet in workbook for row in sheet for cell in row]
        assert "Alpha" in values and "=1+1" in values and "34" in values
        assert all(cell.data_type != "f" for sheet in workbook for row in sheet for cell in row)
    finally:
        workbook.close()
