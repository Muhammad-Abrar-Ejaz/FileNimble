import io
import json
import pytest
from pathlib import Path
from openpyxl import Workbook, load_workbook

from engines.data_converters import (
    csv_to_xlsx,
    xlsx_to_csv,
    json_to_xlsx,
    xlsx_to_json,
    markdown_to_docx,
    docx_to_markdown,
    markdown_to_pdf,
    txt_to_pdf,
    rtf_to_pdf,
)


@pytest.fixture(autouse=True)
def setup_work_dir(tmp_path, monkeypatch):
    monkeypatch.setattr(
        "pathlib.Path", lambda p: tmp_path / Path(p).name if str(p).startswith("/work") else Path(p)
    )


def test_csv_to_xlsx_roundtrip():
    csv_data = "name,score,role\nAlice,95,Engineer\nBob,88,Designer\nCharlie,92,Manager\n".encode(
        "utf-8"
    )
    output, result = csv_to_xlsx(csv_data)
    assert output.exists()
    assert (
        result["mediaType"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert result["rows"] == 4
    assert result["columns"] == 3

    wb = load_workbook(output, data_only=True)
    sheet = wb.active
    assert sheet.cell(1, 1).value == "name"
    assert sheet.cell(2, 1).value == "Alice"
    assert sheet.cell(2, 2).value == "95"
    wb.close()


def test_csv_to_xlsx_formula_injection_defense():
    hostile_csv = "user,payload\nattacker,=cmd|'/C calc'!A0\nnormal,+12345\n".encode("utf-8")
    output, result = csv_to_xlsx(hostile_csv)
    wb = load_workbook(output, data_only=False)
    sheet = wb.active
    # Dangerous formula prefixes are prepended with single quote
    assert sheet.cell(2, 2).value == "'=cmd|'/C calc'!A0"
    assert sheet.cell(3, 2).value == "'+12345"
    # Ensure no cell is classified as formula
    assert all(cell.data_type != "f" for row in sheet.iter_rows() for cell in row)
    wb.close()


def test_csv_to_xlsx_multilingual_unicode():
    unicode_csv = "language,greeting\nJapanese,こんにちは\nArabic,مرحبا\nSpanish,¡Hola!\n".encode(
        "utf-8"
    )
    output, result = csv_to_xlsx(unicode_csv)
    wb = load_workbook(output, data_only=True)
    sheet = wb.active
    assert sheet.cell(2, 2).value == "こんにちは"
    assert sheet.cell(3, 2).value == "مرحبا"
    wb.close()


def test_csv_to_xlsx_null_byte_rejection():
    with pytest.raises(ValueError, match="CORRUPT_FILE"):
        csv_to_xlsx(b"col1,col2\x00malicious")


def test_xlsx_to_csv():
    wb = Workbook()
    ws = wb.active
    ws.append(["id", "name", "city"])
    ws.append([1, "Alice", "New York"])
    ws.append([2, "Bob", "San Francisco"])
    buf = io.BytesIO()
    wb.save(buf)
    wb.close()
    xlsx_bytes = buf.getvalue()

    output, result = xlsx_to_csv(xlsx_bytes)
    assert output.exists()
    assert result["mediaType"] == "text/csv; charset=utf-8"
    content = output.read_text(encoding="utf-8")
    assert "id,name,city" in content
    assert "Alice,New York" in content
    assert "Bob,San Francisco" in content


def test_json_to_xlsx_and_back():
    data = [
        {"product": "Laptop", "price": "1200", "stock": "15"},
        {"product": "Mouse", "price": "25", "stock": "100"},
    ]
    json_bytes = json.dumps(data).encode("utf-8")
    output, result = json_to_xlsx(json_bytes)
    assert output.exists()
    assert (
        result["mediaType"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )

    # Convert back from XLSX to JSON
    xlsx_bytes = output.read_bytes()
    json_output, json_res = xlsx_to_json(xlsx_bytes)
    assert json_output.exists()
    back = json.loads(json_output.read_text(encoding="utf-8"))
    assert len(back) == 2
    assert back[0]["product"] == "Laptop"
    assert back[1]["product"] == "Mouse"


def test_markdown_to_docx_and_back():
    md_content = """# Document Title

## Subsection

This is a paragraph with **bold** text.

- Item A
- Item B

1. First step
2. Second step
"""
    output, result = markdown_to_docx(md_content.encode("utf-8"))
    assert output.exists()
    assert (
        result["mediaType"]
        == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )

    docx_bytes = output.read_bytes()
    md_out, md_res = docx_to_markdown(docx_bytes)
    assert md_out.exists()
    roundtrip = md_out.read_text(encoding="utf-8")
    assert "# Document Title" in roundtrip
    assert "## Subsection" in roundtrip
    assert "**bold**" in roundtrip
    assert "- Item A" in roundtrip


try:
    import weasyprint
except ImportError:
    weasyprint = None


@pytest.mark.skipif(weasyprint is None, reason="weasyprint available in html container")
def test_markdown_to_pdf():
    md = "# Overview\n\nThis is a test of Markdown to PDF conversion.\n\n- Point 1\n- Point 2\n"
    output, result = markdown_to_pdf(md.encode("utf-8"))
    assert output.exists()
    assert result["mediaType"] == "application/pdf"
    assert result["pages"] >= 1
    assert output.read_bytes().startswith(b"%PDF-")


@pytest.mark.skipif(weasyprint is None, reason="weasyprint available in html container")
def test_txt_to_pdf():
    txt = "Plain text report\n=================\nLine 1\nLine 2\nLine 3\n"
    output, result = txt_to_pdf(txt.encode("utf-8"))
    assert output.exists()
    assert result["mediaType"] == "application/pdf"
    assert result["pages"] >= 1
    assert output.read_bytes().startswith(b"%PDF-")


def test_rtf_to_pdf_malformed():
    with pytest.raises(ValueError, match="CORRUPT_FILE"):
        rtf_to_pdf(b"not an rtf document")
